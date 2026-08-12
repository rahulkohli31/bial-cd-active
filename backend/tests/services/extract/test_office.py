"""Office → Markdown extraction (U11) — docx (mammoth) + xlsx (openpyxl), structure gate, caps."""

from __future__ import annotations

import io
from typing import Any

import pytest
from docx import Document
from openpyxl import Workbook

from src.services.extract.office import (
    EXCEL_MEDIA_TYPE,
    MAX_SHEET_ROWS,
    WORD_MEDIA_TYPE,
    OfficeExtractError,
    extract_office,
)


def _docx(*paragraphs: str) -> bytes:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _xlsx(sheet_name: str, rows: list[list[Any]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = sheet_name
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# --- docx ---------------------------------------------------------------------


def test_docx_extracts_paragraph_text() -> None:
    result = extract_office(_docx("Hello world", "Second paragraph"), WORD_MEDIA_TYPE, "doc.docx")
    assert result.format == "word"
    assert "Hello world" in result.text
    assert "Second paragraph" in result.text
    assert result.truncated is False


# --- xlsx ---------------------------------------------------------------------


def test_xlsx_extracts_markdown_table() -> None:
    result = extract_office(
        _xlsx("Data", [["Name", "Age"], ["Alice", 30], ["Bob", 25]]), EXCEL_MEDIA_TYPE, "s.xlsx"
    )
    assert result.format == "excel"
    assert "## Sheet: Data" in result.text
    assert "| Name | Age |" in result.text
    assert "| Alice | 30 |" in result.text


def test_xlsx_escapes_pipes_and_newlines() -> None:
    result = extract_office(_xlsx("S", [["h"], ["a|b\nc"]]), EXCEL_MEDIA_TYPE, "s.xlsx")
    assert "a\\|b c" in result.text  # pipe escaped, newline flattened to space


def test_xlsx_row_cap_truncates() -> None:
    rows: list[list[Any]] = [["h"], *([i] for i in range(MAX_SHEET_ROWS + 50))]
    result = extract_office(_xlsx("Big", rows), EXCEL_MEDIA_TYPE, "big.xlsx")
    assert result.truncated is True
    assert f"showing first {MAX_SHEET_ROWS} rows" in result.text
    assert result.truncation_note != ""


# --- structure gate + errors --------------------------------------------------


def test_missing_zip_signature_rejected() -> None:
    with pytest.raises(OfficeExtractError, match="missing ZIP signature"):
        extract_office(b"plain text, not a zip", WORD_MEDIA_TYPE)


def test_wrong_opc_part_rejected() -> None:
    # A valid xlsx has no word/document.xml → rejected as a Word file.
    xlsx = _xlsx("S", [["a"]])
    with pytest.raises(OfficeExtractError, match="Not a valid Word file"):
        extract_office(xlsx, WORD_MEDIA_TYPE)


def test_unsupported_type_rejected() -> None:
    with pytest.raises(OfficeExtractError, match="Unsupported Office type"):
        extract_office(_xlsx("S", [["a"]]), "application/pdf")
