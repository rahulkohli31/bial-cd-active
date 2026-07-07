"""Parse endpoint (U7, R26): inline + stored modes under the X-App-Key chain, the
route decoded-size cap, and the unsupported-type gate. Storage is an in-memory fake."""

from __future__ import annotations

import base64
import io
from datetime import timedelta

import openpyxl
import pytest
from docx import Document
from sqlalchemy.ext.asyncio import AsyncSession

from src.api.v1.apps.files_router import storage_dependency
from src.db.models.app_registry import AppStatus
from src.services.storage.base import ListPage, ObjectMeta, ObjectStorage
from src.services.storage.errors import StorageNotFoundError
from tests.factories import AppRegistryFactory, UserFactory

_XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class _DictStorage(ObjectStorage):
    def __init__(self) -> None:
        super().__init__(provider="fake")
        self.objects: dict[str, bytes] = {}

    async def put(self, key, data, *, content_type=None, metadata=None):
        self.objects[key] = data
        return ObjectMeta(
            key=key, size=len(data), content_type=content_type, etag=None, last_modified=None
        )

    async def get(self, key):
        if key not in self.objects:
            raise StorageNotFoundError("missing", provider="fake", key=key)
        return self.objects[key]

    async def head(self, key):
        return None

    async def delete(self, key):
        self.objects.pop(key, None)

    async def list(self, prefix, *, page_size=1000, token=None):
        return ListPage(keys=(), next_token=None)

    async def _signed_read_url_impl(self, key, *, expires_in: timedelta):
        return f"https://fake.local/{key}"

    async def aclose(self):
        return None


@pytest.fixture
def store(app) -> _DictStorage:
    fake = _DictStorage()
    app.dependency_overrides[storage_dependency] = lambda: fake
    return fake


async def _approved_app(db: AsyncSession):
    user = await UserFactory.create(db)
    app_row = await AppRegistryFactory.create(
        db, user_id=user.id, status=AppStatus.APPROVED, login_required=False
    )
    return app_row, {"X-App-Key": app_row.app_key}


def _xlsx(rows: list[list[object]]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    for row in rows:
        ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _docx(paragraphs: list[str]) -> bytes:
    doc = Document()
    for para in paragraphs:
        doc.add_paragraph(para)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _b64(data: bytes) -> str:
    return base64.b64encode(data).decode()


# --- inline mode ---------------------------------------------------------------


async def test_inline_xlsx_parses_to_rows(client, store, db_session) -> None:
    app_row, headers = await _approved_app(db_session)
    resp = await client.post(
        f"/v1/apps/{app_row.id}/parse",
        json={
            "filename": "x.xlsx",
            "contentType": _XLSX_TYPE,
            "base64": _b64(_xlsx([["Name"], ["Alice"]])),
        },
        headers=headers,
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["kind"] == "spreadsheet"
    assert body["rows"] == [{"Name": "Alice"}]


async def test_inline_docx_parses_to_text(client, store, db_session) -> None:
    app_row, headers = await _approved_app(db_session)
    resp = await client.post(
        f"/v1/apps/{app_row.id}/parse",
        json={
            "filename": "d.docx",
            "contentType": _DOCX_TYPE,
            "base64": _b64(_docx(["Hello world"])),
        },
        headers=headers,
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["kind"] == "document"
    assert "Hello world" in body["text"]


async def test_inline_unsupported_type_is_415(client, store, db_session) -> None:
    app_row, headers = await _approved_app(db_session)
    resp = await client.post(
        f"/v1/apps/{app_row.id}/parse",
        json={"filename": "x.png", "contentType": "image/png", "base64": _b64(b"\x89PNG")},
        headers=headers,
    )
    assert resp.status_code == 415
    assert resp.json()["error"]["code"] == "UNSUPPORTED_TYPE"


async def test_inline_oversized_is_413(client, store, db_session, monkeypatch) -> None:
    monkeypatch.setattr("src.api.v1.apps.parse_router.APP_FILE_MAX_BYTES", 16)
    app_row, headers = await _approved_app(db_session)
    resp = await client.post(
        f"/v1/apps/{app_row.id}/parse",
        json={"filename": "x.csv", "contentType": "text/csv", "base64": _b64(b"a" * 64)},
        headers=headers,
    )
    assert resp.status_code == 413
    assert resp.json()["error"]["code"] == "FILE_TOO_LARGE"


# --- stored mode ---------------------------------------------------------------


async def test_stored_file_parses(client, store, db_session) -> None:
    app_row, headers = await _approved_app(db_session)
    # Upload a real xlsx via the files endpoint (same fake store), then parse by id.
    up = await client.post(
        f"/v1/apps/{app_row.id}/files",
        json={
            "filename": "s.xlsx",
            "contentType": _XLSX_TYPE,
            "base64": _b64(_xlsx([["N"], ["Zoe"]])),
        },
        headers=headers,
    )
    file_id = up.json()["fileId"]
    resp = await client.post(
        f"/v1/apps/{app_row.id}/parse", json={"fileId": file_id}, headers=headers
    )
    assert resp.status_code == 200
    assert resp.json()["rows"] == [{"N": "Zoe"}]


async def test_stored_unknown_file_is_404(client, store, db_session) -> None:
    from uuid import uuid4

    app_row, headers = await _approved_app(db_session)
    resp = await client.post(
        f"/v1/apps/{app_row.id}/parse", json={"fileId": str(uuid4())}, headers=headers
    )
    assert resp.status_code == 404
